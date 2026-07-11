#!/usr/bin/env python3
"""Export ORIGIN Asset360 Master Data (content + data dictionary) to a styled .xlsx.

Usage:  python3 scripts/export-master.py
Reads DB connection from server/.env (DATABASE_URL). Output: exports/master-data-origin-asset360.xlsx
"""
import subprocess, json, os, re
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "exports", "master-data-origin-asset360.xlsx")
os.makedirs(os.path.dirname(OUT), exist_ok=True)

# Read DATABASE_URL from server/.env (no hardcoded secret)
env_txt = open(os.path.join(ROOT, "server", ".env")).read()
m = re.search(r"^DATABASE_URL=(.+)$", env_txt, re.M)
if not m:
    raise SystemExit("DATABASE_URL not found in server/.env")
DB_URL = m.group(1).strip()


def psql_json(inner_sql):
    sql = f"select coalesce(json_agg(t), '[]'::json) from ({inner_sql}) t"
    out = subprocess.check_output(["psql", DB_URL, "-tAc", sql], text=True)
    return json.loads(out.strip() or "[]")


categories = psql_json("select id, name, created_at from categories order by name")
clients = psql_json("select id, name, created_at from clients order by name")
cols = psql_json("select table_name, ordinal_position, column_name, data_type, is_nullable, column_default "
                 "from information_schema.columns where table_schema='public' order by table_name, ordinal_position")
pks = psql_json("select kcu.table_name, kcu.column_name from information_schema.table_constraints tc "
                "join information_schema.key_column_usage kcu on tc.constraint_name=kcu.constraint_name and tc.table_schema=kcu.table_schema "
                "where tc.constraint_type='PRIMARY KEY' and tc.table_schema='public'")
uniq = psql_json("select kcu.table_name, kcu.column_name from information_schema.table_constraints tc "
                 "join information_schema.key_column_usage kcu on tc.constraint_name=kcu.constraint_name and tc.table_schema=kcu.table_schema "
                 "where tc.constraint_type='UNIQUE' and tc.table_schema='public'")

pk_set = {(r["table_name"], r["column_name"]) for r in pks}
uq_set = {(r["table_name"], r["column_name"]) for r in uniq}

DESC = {
    "categories.id": "ID unik kategori (auto-increment).",
    "categories.name": "Nama kategori aset — UNIK. Dipakai di form pengadaan, filter, & dashboard. Rename ter-cascade ke aset.",
    "categories.created_at": "Waktu data dibuat.",
    "clients.id": "ID unik client (auto-increment).",
    "clients.name": "Nama client/korporasi — UNIK. Dipakai di filter top-bar & form pengadaan. Rename ter-cascade ke aset.",
    "clients.created_at": "Waktu data dibuat.",
    "assets.id": "ID aset (PK, contoh: AST-2026-XXXX).",
    "assets.name": "Nama aset.",
    "assets.category": "Kategori aset — mengacu ke master categories.name.",
    "assets.client": "Client pemilik — mengacu ke master clients.name.",
    "assets.project_code": "Kode Work Order / proyek.",
    "assets.quantity": "Jumlah unit.",
    "assets.current_stage": "Fase siklus hidup saat ini (1–10).",
    "assets.current_location": "Lokasi terkini aset.",
    "assets.qrcode": "Kode QR/label fisik aset.",
    "assets.audit_score": "Skor audit kepatuhan (0–100).",
    "assets.maintenance_status": "Status pemeliharaan: NONE / PENDING / REPAIRING / RESOLVED.",
    "assets.specs": "JSONB — spesifikasi (brand, sku, dimensi, dll).",
    "assets.financials": "JSONB — biaya (purchaseCost, maintenanceCost, disposalValue).",
    "assets.stage_details": "JSONB — detail data tiap fase (request s/d disposal).",
    "assets.created_at": "Waktu aset dibuat.",
    "assets.updated_at": "Waktu terakhir diperbarui.",
    "users.id": "ID unik user.",
    "users.username": "Username login — UNIK.",
    "users.name": "Nama lengkap.",
    "users.password_hash": "Hash password (bcrypt). TIDAK disimpan plaintext.",
    "users.role": "Hak akses: Super Admin / Logistic Supervisor / Field Inspector / Field Technician.",
    "users.created_at": "Waktu akun dibuat.",
    "settings.key": "Kunci setting (PK), mis. company_name, depreciation_pct.",
    "settings.value": "Nilai setting (disimpan sebagai teks).",
    "settings.updated_at": "Waktu terakhir diubah.",
    "activity_logs.id": "ID log (PK).",
    "activity_logs.asset_id": "ID aset terkait.",
    "activity_logs.asset_name": "Nama aset (snapshot).",
    "activity_logs.stage": "Fase saat aksi terjadi.",
    "activity_logs.action": "Deskripsi aktivitas.",
    "activity_logs.operator": "Pelaku aksi.",
    "activity_logs.type": "Jenis: info / success / warning / error.",
    "activity_logs.timestamp": "Waktu aktivitas.",
}
TABLE_DESC = {
    "categories": "Master Data — daftar kategori aset.",
    "clients": "Master Data — daftar client/korporasi.",
    "assets": "Data utama aset & seluruh siklus hidupnya.",
    "users": "Akun operator & hak akses (role).",
    "settings": "Konfigurasi sistem (key-value).",
    "activity_logs": "Jejak audit aktivitas sistem.",
}

NAVY, DARK = "1D4ED8", "0A0F1D"
thin = Side(style="thin", color="D9E1F2")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
HFONT = Font(bold=True, color="FFFFFF", size=11)
HFILL = PatternFill("solid", fgColor=NAVY)
TITLEF = Font(bold=True, color="FFFFFF", size=14)
TITLEFILL = PatternFill("solid", fgColor=DARK)


def fmt_dt(v):
    if not v:
        return ""
    try:
        return datetime.fromisoformat(str(v)).strftime("%d %b %Y %H:%M")
    except Exception:
        return str(v)


def write_sheet(ws, title, headers, rows, subtitle=""):
    ncol = len(headers)
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncol)
    c = ws.cell(1, 1, title)
    c.font, c.fill, c.alignment = TITLEF, TITLEFILL, Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 30
    start = 2
    if subtitle:
        ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=ncol)
        s = ws.cell(2, 1, subtitle)
        s.font = Font(italic=True, color="64748B", size=9)
        s.alignment = Alignment(indent=1)
        start = 3
    hr = start
    for j, h in enumerate(headers, 1):
        cell = ws.cell(hr, j, h)
        cell.font, cell.fill = HFONT, HFILL
        cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)
        cell.border = BORDER
    ws.row_dimensions[hr].height = 22
    for i, row in enumerate(rows):
        for j, val in enumerate(row, 1):
            cell = ws.cell(hr + 1 + i, j, val)
            cell.border = BORDER
            cell.alignment = Alignment(vertical="top", horizontal="left", indent=1, wrap_text=(len(str(val)) > 40))
            if (i % 2) == 1:
                cell.fill = PatternFill("solid", fgColor="F8FAFC")
    for j, h in enumerate(headers, 1):
        maxlen = len(str(h))
        for row in rows:
            maxlen = max(maxlen, len(str(row[j - 1])))
        ws.column_dimensions[get_column_letter(j)].width = min(max(maxlen + 3, 10), 60)
    ws.freeze_panes = ws.cell(hr + 1, 1)
    ws.sheet_view.showGridLines = False


def dict_rows(table_names):
    rows = []
    for t in table_names:
        for c in [x for x in cols if x["table_name"] == t]:
            key = f"{t}.{c['column_name']}"
            keys = []
            if (t, c["column_name"]) in pk_set:
                keys.append("PK")
            if (t, c["column_name"]) in uq_set:
                keys.append("UNIQUE")
            rows.append([t, c["column_name"], c["data_type"], "Tidak" if c["is_nullable"] == "YES" else "Ya",
                         (c["column_default"] or ""), ", ".join(keys), DESC.get(key, "")])
    return rows


wb = Workbook()
ws = wb.active
ws.title = "Ringkasan"
write_sheet(ws, "ORIGIN Asset360 — Export Master Data", ["Field", "Keterangan"], [
    ["Dokumen", "Export Master Data — ORIGIN Asset360"],
    ["Dibuat", datetime.now().strftime("%d %B %Y %H:%M WIB")],
    ["Database", "origin_asset360 (PostgreSQL 16)"],
    ["Total Kategori", str(len(categories))],
    ["Total Client", str(len(clients))],
    ["Catatan", "Sheet 'Kategori Aset' & 'Client' = isi data. Sheet 'Kamus Data' = struktur tabel & kolom untuk penyesuaian pra go-live."],
])
ws = wb.create_sheet("Kategori Aset")
write_sheet(ws, "Master Data — Kategori Aset", ["ID", "Nama Kategori", "Dibuat"],
            [[r["id"], r["name"], fmt_dt(r["created_at"])] for r in categories],
            subtitle=f"{len(categories)} kategori · tabel: categories")
ws = wb.create_sheet("Client")
write_sheet(ws, "Master Data — Client / Korporasi", ["ID", "Nama Client", "Dibuat"],
            [[r["id"], r["name"], fmt_dt(r["created_at"])] for r in clients],
            subtitle=f"{len(clients)} client · tabel: clients")
DICT_HEADERS = ["Tabel", "Kolom", "Tipe Data", "Wajib Diisi", "Nilai Default", "Kunci", "Keterangan"]
ws = wb.create_sheet("Kamus Data - Master")
write_sheet(ws, "Kamus Data — Tabel Master (categories, clients)", DICT_HEADERS,
            dict_rows(["categories", "clients"]), subtitle="Struktur kolom tabel Master Data")
allt = ["categories", "clients", "assets", "users", "settings", "activity_logs"]
ws = wb.create_sheet("Kamus Data - Semua Tabel")
write_sheet(ws, "Kamus Data — Seluruh Skema (referensi go-live)", DICT_HEADERS, dict_rows(allt),
            subtitle="Semua tabel di database. " + " | ".join(f"{t}: {TABLE_DESC.get(t,'')}" for t in allt))

wb.save(OUT)
print("SAVED:", OUT)
print("sheets:", wb.sheetnames, "| categories:", len(categories), "clients:", len(clients))
